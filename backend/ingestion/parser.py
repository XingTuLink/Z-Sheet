"""XLSX/CSV parsing into ParsedWorkbook (Day 5).

Stateless by design: bytes in, structured rows out, no database or file
storage. Cells are normalized uniformly to trimmed strings (None when blank)
so later stages never branch on the source format.

Defensive V0.1 limits reject pathological inputs explicitly rather than
silently truncating them, because truncated rows would corrupt the profiling
and type-inference stages that follow.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend.logging_config import get_logger

from .formulas import evaluate_formulas
from .header_detection import detect_header_row, used_width
from .schemas import Cell, ParsedSheet, ParsedWorkbook

logger = get_logger(__name__)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_DATA_ROWS = 50_000
MAX_COLUMNS = 200
# Hard guard against worksheets whose declared dimension points at the full
# 1,048,576-row sheet (read_only iteration would otherwise walk blanks).
MAX_RAW_ROWS = 200_000

_ZIP_MAGIC = b"PK\x03\x04"
# NOTE: bare "utf-16" must never sit in a probe chain without a BOM gate —
# without a BOM the codec assumes native byte order and decodes arbitrary
# even-length byte streams (e.g. GBK) into garbage without raising.
_UTF16_BOM = (b"\xff\xfe", b"\xfe\xff")
_CSV_ENCODINGS = ("utf-8-sig", "gb18030")
_DELIMITERS = (",", "\t", ";")


class ParseError(ValueError):
    """A user-facing input problem; mapped to HTTP 422 by the route."""


def parse_workbook(file_name: str, content: bytes) -> ParsedWorkbook:
    if not content:
        raise ParseError("uploaded file is empty")
    file_type = _detect_file_type(file_name, content)
    if file_type == "xlsx":
        sheets = _parse_xlsx(content)
    else:
        sheets = [_parse_csv(file_name, content)]
    return ParsedWorkbook(file_name=file_name, file_type=file_type, sheets=sheets)


def _detect_file_type(file_name: str, content: bytes) -> Literal["xlsx", "csv"]:
    lower = file_name.lower()
    if content.startswith(_ZIP_MAGIC):
        # openpyxl performs the real validation below; a renamed .docx/.zip
        # reaches the same error path.
        return "xlsx"
    if lower.endswith(".xlsx"):
        raise ParseError("file has an .xlsx name but is not an XLSX container")
    if lower.endswith(".xls"):
        raise ParseError("legacy .xls workbooks are not supported; save the file as .xlsx")
    return "csv"


# ---------- XLSX ----------


def _cell_to_text(value: object) -> Cell:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _parse_xlsx(content: bytes) -> list[ParsedSheet]:
    try:
        # Normal (non read-only) mode is required for merged-cell ranges:
        # real-world Excel workbooks encode categories as vertical merges
        # ("AI 智能体开发" once over six service rows), and only the anchor
        # cell carries the value in that layout.
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ParseError(f"not a valid XLSX file: {exc}") from exc

    # A formula cell whose result was never cached reads back as None under
    # data_only=True (silently emptying a whole column when a writer never
    # computed it). A cheap read-only pass locates those cells so they can be
    # recalculated or reported instead of dropping the data.
    formulas = _formula_cell_texts(content)

    prepared: list[tuple[str, list[list[Cell]], dict[int, list[tuple[int, int, str]]]]] = []
    try:
        for worksheet in workbook.worksheets:
            # Vertical merges only (single column spanning rows): forward-fill
            # the anchor value into the covered rows. Horizontal merges
            # (titles, annotation/contact banners, two-row group headers) stay
            # untouched so the blank/annotation row filters keep treating them
            # as non-data; their spans are recorded for grouped-header merging.
            vertical_fills: list[tuple[int, int, int, str, str]] = []
            horizontal_spans: dict[int, list[tuple[int, int, str]]] = {}
            for merged in worksheet.merged_cells.ranges:
                if merged.min_row == merged.max_row and merged.min_col < merged.max_col:
                    group_label = _cell_to_text(
                        worksheet.cell(merged.min_row, merged.min_col).value
                    )
                    if group_label:
                        horizontal_spans.setdefault(merged.min_row - 1, []).append(
                            (merged.min_col, merged.max_col, group_label)
                        )
                    continue
                if merged.min_col != merged.max_col or merged.min_row >= merged.max_row:
                    continue
                raw_anchor = _cell_to_text(
                    worksheet.cell(merged.min_row, merged.min_col).value
                )
                if raw_anchor:
                    # In-cell line breaks inside a merged category label are
                    # purely visual wrapping ("AI 智能体\\n开发"); collapse them
                    # so enum profiling sees one consistent value across the
                    # anchor and the filled rows.
                    anchor = re.sub(r"\s*[\r\n]+\s*", "", raw_anchor).strip()
                    if anchor:
                        vertical_fills.append(
                            (
                                merged.min_col,
                                merged.min_row,
                                merged.max_row,
                                raw_anchor,
                                anchor,
                            )
                        )

            rows: list[list[Cell]] = []
            for row in worksheet.iter_rows(values_only=True):
                rows.append([_cell_to_text(value) for value in row])
                if len(rows) > MAX_RAW_ROWS:
                    raise ParseError(
                        f"sheet {worksheet.title!r} exceeds {MAX_RAW_ROWS} raw rows"
                    )

            # Coordinates are 1-based; rows list is 0-based. Normalize the
            # anchor cell itself, then fill empty covered cells, never
            # overwriting explicit values.
            for col, anchor_row, row_end, raw_anchor, anchor in vertical_fills:
                col_idx = col - 1
                anchor_idx = anchor_row - 1
                if 0 <= anchor_idx < len(rows) and col_idx < len(rows[anchor_idx]):
                    if rows[anchor_idx][col_idx] == raw_anchor:
                        rows[anchor_idx][col_idx] = anchor
                for row_no in range(anchor_row + 1, row_end + 1):
                    row_idx = row_no - 1
                    if not (0 <= row_idx < len(rows)):
                        continue
                    cells = rows[row_idx]
                    if col_idx >= len(cells):
                        cells.extend([None] * (col_idx + 1 - len(cells)))
                    if cells[col_idx] is None:
                        cells[col_idx] = anchor

            prepared.append((worksheet.title, rows, horizontal_spans))
    finally:
        workbook.close()

    # Recalculation needs every sheet's grid, because formulas reference
    # sibling sheets (SUMIF over 订单明细 from 订单).
    formula_counts = evaluate_formulas(
        {title: rows for title, rows, _ in prepared}, formulas
    )
    sheets: list[ParsedSheet] = []
    for title, rows, horizontal_spans in prepared:
        resolved, missing = formula_counts.get(title, (0, 0))
        sheets.append(
            _build_sheet(
                title,
                rows,
                horizontal_spans=horizontal_spans,
                formula_resolved=resolved,
                formula_missing=missing,
            )
        )
    return sheets


def _formula_cell_texts(content: bytes) -> dict[str, dict[tuple[int, int], str]]:
    """0-based (row, col) -> formula text, per sheet — best-effort.

    A formula whose result was never cached reads back as None under
    ``data_only=True``, silently emptying the cell. The scan is advisory:
    any failure simply degrades to "no formulas detected" rather than
    failing the whole parse — but it must leave a trace, otherwise a real
    scan failure is indistinguishable from a genuinely formula-free file.
    """
    try:
        book = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    except Exception:
        logger.warning("公式单元格扫描：工作簿无法只读打开，本次跳过公式列告警", exc_info=True)
        return {}
    try:
        texts: dict[str, dict[tuple[int, int], str]] = {}
        for worksheet in book.worksheets:
            found: dict[tuple[int, int], str] = {}
            for row_index, row in enumerate(worksheet.iter_rows()):
                for col_index, cell in enumerate(row):
                    if cell.data_type == "f" and isinstance(cell.value, str):
                        found[(row_index, col_index)] = cell.value
            texts[worksheet.title] = found
        return texts
    except Exception:
        logger.warning("公式单元格扫描：遍历工作表失败，本次跳过公式列告警", exc_info=True)
        return {}
    finally:
        book.close()


def _normalize_header(value: str) -> str:
    """Normalize a header label: full-width → ASCII (NFKC), then trim."""
    return unicodedata.normalize("NFKC", value).strip()


def _apply_group_headers(
    columns: list[str], spans: list[tuple[int, int, str]]
) -> list[str]:
    """Prefix child column names with their two-row group label.

    ``spans`` are 1-based ``(min_col, max_col, label)`` horizontal merges on
    the row above the header. Unnamed placeholder columns (``column_n``) are
    never prefixed, so a blank sub-header does not become ``组名_column_n``.
    """
    merged = list(columns)
    for min_col, max_col, label in spans:
        group = _normalize_header(label)
        if not group:
            continue
        stop = min(max_col, len(merged))
        for pos in range(min_col - 1, stop):
            child = merged[pos]
            if child.startswith("column_"):
                continue
            merged[pos] = f"{group}_{child}"
    return merged


# ---------- CSV ----------


def _decode_csv(content: bytes) -> str:
    # Excel on Chinese Windows exports GBK/GB18030; "CSV UTF-8" carries a BOM
    # (handled by utf-8-sig); Excel's "Unicode Text" CSV is UTF-16 with BOM.
    if content.startswith(_UTF16_BOM):
        candidates = ("utf-16", "utf-8-sig", "gb18030")
    else:
        candidates = _CSV_ENCODINGS
    last_error: UnicodeDecodeError | None = None
    for encoding in candidates:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ParseError(
        "cannot decode CSV as UTF-8 or GB18030; re-export it as UTF-8 encoded CSV"
    ) from last_error


def _detect_delimiter(text: str) -> str:
    counts = {delimiter: 0 for delimiter in _DELIMITERS}
    checked = 0
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        for delimiter in _DELIMITERS:
            counts[delimiter] += stripped.count(delimiter)
        checked += 1
        if checked >= 10:
            break
    return max(_DELIMITERS, key=lambda delimiter: counts[delimiter])


def _parse_csv(file_name: str, content: bytes) -> ParsedSheet:
    text = _decode_csv(content)
    delimiter = _detect_delimiter(text)
    rows: list[list[Cell]] = []
    for record in csv.reader(io.StringIO(text), delimiter=delimiter):
        if not record:
            rows.append([])
            continue
        rows.append([(cell.strip() or None) for cell in record])
        if len(rows) > MAX_RAW_ROWS:
            raise ParseError(f"CSV exceeds {MAX_RAW_ROWS} raw rows")
    sheet_name = Path(file_name).stem or "sheet1"
    return _build_sheet(sheet_name, rows)


# ---------- Shared sheet assembly ----------


# A row is blank when it carries no "content character": no letter, digit or
# CJK ideograph. Rows made only of ".", "-", "/", spaces or punctuation are
# template artifacts (a stray "." in the remark column otherwise doubles the
# row count and destroys uniqueness statistics on small sheets).
_CONTENT_CHAR_RE = re.compile(r"[0-9A-Za-z\u4e00-\u9fff]")

# Aggregate/footer rows ("合计 3138") are spreadsheet totals, not business
# records. Their per-row key column is empty, so they otherwise fail key
# validation at Confirm time.
_TOTAL_MARKERS = frozenset(
    {"合计", "共计", "总计", "小计", "累计", "total", "sum", "grand total"}
)

# Running-number column headers. When such a column holds an established
# integer sequence, every later row whose cell there is *not* a number is
# material pasted under the table (special notes, contact blocks, banners)
# rather than a record — that is exactly the junk that otherwise crashes type
# coercion ("'💡 特别说明' is not a number") and blocks generation.
_SEQUENCE_HEADERS = frozenset(
    {
        "序号", "序", "序號", "编号", "編號", "编码", "行号", "行號", "项次",
        "no", "no.", "#", "num", "number", "index", "idx",
    }
)
# Pure integers only ("34" or Excel's "34.0"). Business codes like "A001"
# never match, so a 编号 column carrying codes disables the rule entirely.
_INTEGER_CELL_RE = re.compile(r"^[+-]?\d+(?:\.0+)?$")
# Records are numbered from row 1; a leading run this long (blank
# continuation cells skipped) proves the column really is a running sequence,
# regardless of how much note/footer junk follows it.
_SEQUENCE_MIN_RUN = 3

# Annotation rows ("💡 特别说明", "✅ 我们可以……") are single-column notes.
_LEADING_EMOJI_RE = re.compile(
    r"^[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF]"
)
_NOTE_PREFIXES = (
    "注：", "注:", "备注：", "备注:", "说明：", "说明:",
    "特别说明", "温馨提示", "小贴士", "注意", "※",
)


def _is_blank(row: list[Cell]) -> bool:
    return not any(
        cell is not None and _CONTENT_CHAR_RE.search(cell) for cell in row
    )


def _is_summary_row(row: list[Cell]) -> bool:
    for cell in row:
        if cell is None:
            continue
        text = cell.strip().strip(":：").strip().lower()
        if text in _TOTAL_MARKERS:
            return True
    return False


def _detect_sequence_column(
    columns: list[str], rows: list[list[Cell]]
) -> int | None:
    """Index of a running-number column, or None.

    Enabled only when the header is sequence-like AND the first populated
    cells below it form an integer run. Junk/footer rows that come *after*
    the records must not depress the evidence (chicken-and-egg), and blank
    continuation cells of a merged layout are skipped.
    """
    if not columns or columns[0].strip().lower() not in _SEQUENCE_HEADERS:
        return None
    run = 0
    for row in rows:
        if not row or row[0] is None:
            continue
        if _INTEGER_CELL_RE.match(row[0].strip()):
            run += 1
            continue
        break
    return 0 if run >= _SEQUENCE_MIN_RUN else None


def _is_annotation_row(row: list[Cell]) -> bool:
    """A note/banner line: content sits in one cell and looks like a note."""
    populated = [cell for cell in row if cell is not None]
    if len(populated) != 1:
        return False
    text = populated[0].strip()
    if _LEADING_EMOJI_RE.match(text):
        return True
    return any(text.startswith(prefix) for prefix in _NOTE_PREFIXES)


def _build_sheet(
    name: str,
    rows: list[list[Cell]],
    *,
    horizontal_spans: dict[int, list[tuple[int, int, str]]] | None = None,
    formula_resolved: int = 0,
    formula_missing: int = 0,
) -> ParsedSheet:
    detection = detect_header_row(rows)
    if detection is None:
        return ParsedSheet(name=name, is_empty=True)

    width = used_width(rows)
    if width > MAX_COLUMNS:
        raise ParseError(f"sheet {name!r} has {width} columns; limit is {MAX_COLUMNS}")

    warnings: list[str] = []
    if formula_missing or formula_resolved:
        if formula_missing:
            calculated = (
                f"{formula_resolved} 个已按公式自动计算；" if formula_resolved else ""
            )
            warnings.append(
                f"有 {formula_missing} 个公式单元格无法计算（{calculated}"
                "源文件没有缓存计算结果），已按空值读取；"
                "请用 Excel 打开后重新保存（或粘贴计算后的数值）再上传"
            )
        else:
            warnings.append(
                f"有 {formula_resolved} 个公式单元格没有缓存计算结果，已按公式自动计算"
            )
    if detection.confidence == "low":
        warnings.append(
            f"第 {detection.row_index + 1} 行被识别为表头，但把握不大，请核对"
        )

    raw_header = rows[detection.row_index]
    columns: list[str] = []
    unnamed = 0
    for pos in range(width):
        value = raw_header[pos] if pos < len(raw_header) else None
        header = _normalize_header(value) if value is not None else ""
        if header:
            columns.append(header)
        else:
            unnamed += 1
            columns.append(f"column_{pos + 1}")
    if unnamed:
        warnings.append(f"有 {unnamed} 个表头没有名称，已按 column_<n> 生成")

    # A two-row grouped header (e.g. 收款信息 / 退款信息 each containing a
    # 金额 column) is disambiguated by prefixing child names with the group
    # label. Only a genuine group row qualifies: two or more horizontal
    # merges directly above the header; a single merged banner does not.
    group_spans = (horizontal_spans or {}).get(detection.row_index - 1, [])
    if len(group_spans) >= 2:
        columns = _apply_group_headers(columns, group_spans)

    duplicates = sorted({col for col in columns if columns.count(col) > 1})
    if duplicates:
        warnings.append(f"表头存在重复：{', '.join(duplicates)}")

    # First pass: structural cleanup only (blank/summary rows), aligned to
    # the detected width. The sequence-column probe needs these candidates to
    # decide whether the number column is a real running sequence.
    candidates: list[list[Cell]] = []
    summary_rows_dropped = 0
    for raw in rows[detection.row_index + 1 :]:
        if _is_blank(raw):
            continue
        aligned = [raw[pos] if pos < len(raw) else None for pos in range(width)]
        if _is_summary_row(aligned):
            summary_rows_dropped += 1
            continue
        candidates.append(aligned)

    sequence_col = _detect_sequence_column(columns, candidates)

    # Second pass: drop content pasted under the table — notes, banners and
    # footer blocks — so it can never block record extraction/Confirm.
    data_rows: list[list[Cell]] = []
    annotation_rows_dropped = 0
    annotation_samples: list[str] = []
    for aligned in candidates:
        sequence_cell = (
            aligned[sequence_col] if sequence_col is not None else None
        )
        breaks_sequence = (
            sequence_cell is not None
            and not _INTEGER_CELL_RE.match(sequence_cell.strip())
        )
        if breaks_sequence or _is_annotation_row(aligned):
            annotation_rows_dropped += 1
            if len(annotation_samples) < 2:
                first_filled = next(
                    (cell for cell in aligned if cell is not None), ""
                )
                annotation_samples.append(first_filled[:20])
            continue
        data_rows.append(aligned)
        if len(data_rows) > MAX_DATA_ROWS:
            raise ParseError(
                f"sheet {name!r} exceeds {MAX_DATA_ROWS} data rows; V0.1 limit"
            )

    if summary_rows_dropped:
        warnings.append(
            f"已忽略 {summary_rows_dropped} 行合计/总计/小计类汇总行"
        )
    if annotation_rows_dropped:
        sample = f"，例如 {annotation_samples!r}" if annotation_samples else ""
        if sequence_col is not None:
            warnings.append(
                f"已忽略 {annotation_rows_dropped} 行不属于「{columns[0]}」序号序列的"
                f"非记录行（表尾注释或联系块{sample}）；如果这些是真实数据，"
                "请从工作表中移除"
            )
        else:
            warnings.append(
                f"已忽略 {annotation_rows_dropped} 行注释或横幅内容{sample}"
            )

    if not data_rows:
        warnings.append("工作表有表头但没有数据行")

    return ParsedSheet(
        name=name,
        header_row_number=detection.row_index + 1,
        columns=columns,
        rows=data_rows,
        data_row_count=len(data_rows),
        warnings=warnings,
    )
