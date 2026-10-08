"""Tests for Day 10 deterministic entity recognition."""

from __future__ import annotations

import io
from collections.abc import Sequence

import openpyxl
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.ingestion.schemas import ParsedSheet, ParsedWorkbook
from backend.understanding.profiling import add_profiles
from backend.understanding.schemas import InferredEntity

client = TestClient(app)

ORDER_ROWS = [
    ["SO-1001", "杭州云栖", "87350.5", "2026-07-03"],
    ["SO-1002", "杭州云栖", "12000", "2026-08-15"],
    ["SO-1003", "深圳前海", "6400", "2026-09-01"],
    ["SO-1004", "成都锦江", "3200", "2026-09-20"],
    ["SO-1005", "北京海淀", "9800", "2026-10-02"],
    ["SO-1006", "南京玄武", "15600", "2026-10-05"],
]
ORDER_COLUMNS = ["订单编号", "客户名称", "金额", "下单日期"]

CUSTOMER_ROWS = [
    ["C001", "杭州云栖科技", "华东", "13800001111"],
    ["C002", "深圳前海贸易", "华南", "13900002222"],
    ["C003", "成都锦江商贸", "华东", "13700003333"],
    ["C004", "北京海淀贸易", "华北", "13600004444"],
    ["C005", "南京玄武商贸", "华东", "13500005555"],
    ["C006", "武汉江汉贸易", "华南", "13400006666"],
]
CUSTOMER_COLUMNS = ["客户编号", "客户名称", "区域", "联系电话"]

PRODUCT_ROWS = [
    ["P001", "无线鼠标", "￥99.00", "黑色", "120"],
    ["P002", "机械键盘", "￥299.00", "黑色", "45"],
    ["P003", "显示器支架", "￥159.00", "银色", "8"],
    ["P004", "USB-C 线材", "￥29.90", "白色", "300"],
    ["P005", "笔记本内胆包", "￥49.00", "灰色", "60"],
    ["P006", "无线鼠标", "￥99.00", "白色", "0"],
]
PRODUCT_COLUMNS = ["商品编号", "商品名称", "单价", "规格", "库存数量"]


def _sheet(name: str, columns: list[str], rows: Sequence[Sequence[object]]) -> ParsedSheet:
    normalized = [[None if cell is None else str(cell) for cell in row] for row in rows]
    return ParsedSheet(
        name=name,
        header_row_number=1,
        columns=columns,
        rows=normalized,
        data_row_count=len(normalized),
    )


def entity_for(
    name: str, columns: list[str], rows: Sequence[Sequence[object]]
) -> InferredEntity:
    workbook = ParsedWorkbook(
        file_name="f.xlsx",
        file_type="xlsx",
        sheets=[_sheet(name, columns, rows)],
    )
    entity = add_profiles(workbook).sheets[0].inferred_entity
    assert entity is not None
    return entity


# ---------- sheet-name evidence, structure confirmed ----------


def test_named_order_sheet():
    entity = entity_for("订单", ORDER_COLUMNS, ORDER_ROWS)
    assert entity.key == "order" and entity.name == "订单"
    assert entity.confidence == 0.95 and entity.needs_review is False
    assert entity.key_field == "订单编号"
    assert "structure_confirmed" in entity.signals


def test_named_customer_sheet():
    entity = entity_for("客户", CUSTOMER_COLUMNS, CUSTOMER_ROWS)
    assert entity.key == "customer" and entity.name == "客户"
    assert entity.confidence >= 0.85 and entity.needs_review is False
    assert entity.key_field == "客户编号"


def test_named_product_sheet():
    entity = entity_for("商品", PRODUCT_COLUMNS, PRODUCT_ROWS)
    assert entity.key == "product" and entity.confidence >= 0.85
    assert entity.key_field == "商品编号"
    assert "has_price_field" in entity.signals


def test_sheet_name_table_suffix_is_stripped():
    entity = entity_for("订单表", ORDER_COLUMNS, ORDER_ROWS)
    assert entity.key == "order" and entity.name == "订单"


# ---------- default names: structure carries the decision ----------


def test_default_named_sheets_classified_by_structure():
    order = entity_for("Sheet1", ORDER_COLUMNS, ORDER_ROWS)
    assert order.key == "order" and order.confidence == 0.85
    assert order.needs_review is False
    assert order.name == "订单"

    customer = entity_for("Sheet2", CUSTOMER_COLUMNS, CUSTOMER_ROWS)
    assert customer.key == "customer" and customer.confidence == 0.85

    product = entity_for("Sheet3", PRODUCT_COLUMNS, PRODUCT_ROWS)
    assert product.key == "product" and product.confidence == 0.85


def test_mid_order_shape_without_money_stays_medium():
    rows = [[row[0], row[1], row[3]] for row in ORDER_ROWS]
    entity = entity_for("Sheet4", ["订单编号", "客户名称", "下单日期"], rows)
    assert entity.key == "order" and entity.confidence == 0.75
    assert entity.needs_review is True
    assert "partial_transaction_shape" in entity.signals


def test_party_name_only_is_weak_customer():
    rows = [["E001", "张伟"], ["E002", "李娜"], ["E003", "王强"], ["E004", "赵敏"]]
    entity = entity_for("Sheet5", ["编号", "名称"], rows)
    assert entity.key == "customer" and entity.confidence == 0.70
    assert entity.needs_review is True
    assert "party_name_only" in entity.signals


# ---------- conflict / missing key / unknown ----------


def test_sheet_name_conflicts_with_strong_structure():
    # A sheet labelled 客户 but holding the order triad.
    entity = entity_for("客户", ORDER_COLUMNS, ORDER_ROWS)
    assert entity.key == "customer"  # name direction kept
    assert entity.confidence == 0.61 and entity.needs_review is True
    assert "name_structure_conflict" in entity.signals
    assert "structure_suggests:order" in entity.signals


def test_order_sheet_without_identifier_is_flagged():
    columns = ["客户名称", "金额", "下单日期"]
    rows = [row[1:] for row in ORDER_ROWS]
    entity = entity_for("订单", columns, rows)
    assert entity.key == "order"
    assert entity.key_field is None
    assert entity.confidence == 0.75 and entity.needs_review is True
    assert "no_identifier" in entity.signals


def test_lookup_table_is_unknown_not_forced_into_a_kind():
    rows = [
        ["系统名称", "销售管理"],
        ["会话超时", "30"],
        ["总部城市", "上海"],
        ["启用审批", "是"],
        ["启用短信", "否"],
        ["实施顾问", "张三"],
    ]
    entity = entity_for("Sheet9", ["参数名", "参数值"], rows)
    assert entity.key == "unknown"
    assert entity.confidence == 0.5 and entity.needs_review is True
    assert entity.name == "未命名实体"


def test_empty_sheet_has_no_entity():
    empty = ParsedSheet(
        name="空表",
        is_empty=True,
        header_row_number=None,
        columns=[],
        rows=[],
        data_row_count=0,
    )
    workbook = ParsedWorkbook(file_name="f.xlsx", file_type="xlsx", sheets=[empty])
    sheet = add_profiles(workbook).sheets[0]
    assert sheet.is_empty is True
    assert sheet.inferred_entity is None
    assert sheet.inferred_fields == []


# ---------- end to end: real multi-sheet xlsx ----------


def test_parse_endpoint_recognizes_three_entities():
    wb = openpyxl.Workbook()
    wb.remove(wb.worksheets[0])
    for name, columns, rows in (
        ("订单", ORDER_COLUMNS, ORDER_ROWS),
        ("客户", CUSTOMER_COLUMNS, CUSTOMER_ROWS),
        ("商品", PRODUCT_COLUMNS, PRODUCT_ROWS),
    ):
        ws = wb.create_sheet(name)
        ws.append(columns)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)

    response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("ledger.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    entities = {
        sheet["name"]: sheet["inferred_entity"]
        for sheet in response.json()["sheets"]
    }
    assert entities["订单"]["key"] == "order"
    assert entities["客户"]["key"] == "customer"
    assert entities["商品"]["key"] == "product"
    assert entities["订单"]["key_field"] == "订单编号"
