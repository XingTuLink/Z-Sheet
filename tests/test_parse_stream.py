"""Tests for Day 12: SSE parse stream with real stage progress events."""

from __future__ import annotations

import io
import json

import openpyxl
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.ingestion.parser import parse_workbook
from backend.ingestion.schemas import ParsedWorkbook
from backend.understanding.profiling import (
    STAGE_ASSEMBLE,
    STAGE_ENTITIES,
    STAGE_FIELDS,
    STAGE_RELATIONS,
    STAGE_SHEETS,
    add_profiles,
)

client = TestClient(app)


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.title = "客户"
    ws.append(["客户名称", "区域", "联系电话"])
    ws.append(["杭州云栖", "华东", "13800000001"])
    ws.append(["深圳前海", "华南", "13800000002"])

    ws2 = wb.create_sheet("订单")
    ws2.append(["订单编号", "客户名称", "金额", "下单日期"])
    ws2.append(["SO-1", "杭州云栖", 100, "2026-01-01"])
    ws2.append(["SO-2", "深圳前海", 200, "2026-02-01"])
    ws2.append(["SO-3", "杭州云栖", 300, "2026-03-01"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _frames(text: str) -> list[tuple[str, str]]:
    frames: list[tuple[str, str]] = []
    for block in text.strip().split("\n\n"):
        event = ""
        data = ""
        for line in block.splitlines():
            if line.startswith("event: "):
                event = line[len("event: "):]
            elif line.startswith("data: "):
                data = line[len("data: "):]
        frames.append((event, data))
    return frames


def test_stream_emits_real_stages_then_result():
    resp = client.post(
        "/api/v1/ingestion/parse/stream",
        files={"file": ("台账.xlsx", _xlsx_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/event-stream")

    frames = _frames(resp.text)
    kinds = [event for event, _ in frames]
    assert kinds == ["stage", "stage", "stage", "stage", "stage", "result"]

    stages = [(json.loads(data)["key"], json.loads(data)["count"])
              for event, data in frames[:-1]]
    assert [key for key, _ in stages] == [
        STAGE_SHEETS, STAGE_FIELDS, STAGE_ENTITIES, STAGE_RELATIONS, STAGE_ASSEMBLE,
    ]
    counts = dict(stages)
    assert counts[STAGE_SHEETS] == 2
    assert counts[STAGE_FIELDS] == 7  # 3 customer columns + 4 order columns
    assert counts[STAGE_ENTITIES] == 2
    assert counts[STAGE_RELATIONS] == 1
    assert counts[STAGE_ASSEMBLE] == 0

    result = json.loads(frames[-1][1])
    assert result["business_model"] is not None
    assert {e["key"] for e in result["business_model"]["entities"]} == {
        "customer", "order",
    }
    assert result["business_model_yaml"]


def test_stream_reports_parse_error_as_event():
    resp = client.post(
        "/api/v1/ingestion/parse/stream",
        files={"file": ("broken.xlsx", b"this is not a zip/xlsx file",
                        "application/octet-stream")},
    )
    assert resp.status_code == 200  # errors travel inside the stream
    frames = _frames(resp.text)
    assert [event for event, _ in frames] == ["error"]
    detail = json.loads(frames[0][1])["detail"]
    assert isinstance(detail, str) and detail


def test_add_profiles_stage_callback_order_and_counts():
    workbook = parse_workbook("台账.xlsx", _xlsx_bytes())
    seen: list[tuple[str, int]] = []
    add_profiles(workbook, on_stage=lambda key, count: seen.append((key, count)))
    assert [key for key, _ in seen] == [
        STAGE_FIELDS, STAGE_ENTITIES, STAGE_RELATIONS, STAGE_ASSEMBLE,
    ]
    assert dict(seen)[STAGE_FIELDS] == 7
    assert dict(seen)[STAGE_ENTITIES] == 2
    assert dict(seen)[STAGE_RELATIONS] == 1


def test_add_profiles_without_callback_still_works():
    # Explicit ParsedWorkbook so the test does not depend on the parser:
    # the optional argument must stay backward compatible.
    from backend.ingestion.schemas import ParsedSheet

    empty = ParsedSheet(name="空表", is_empty=True, header_row_number=None,
                        columns=[], rows=[])
    wb = ParsedWorkbook(file_name="empty.xlsx", file_type="xlsx", sheets=[empty])
    result = add_profiles(wb)
    assert result.business_model is None
