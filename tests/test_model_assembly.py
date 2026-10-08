"""Tests for Day 11: one-to-many link inference and Business Model assembly."""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import TypeVar

import openpyxl
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.domain.models import (
    BusinessModel,
    DashboardView,
    DetailView,
    ListView,
)
from backend.domain.serialization import model_from_yaml
from backend.ingestion.schemas import ParsedSheet, ParsedWorkbook
from backend.understanding.profiling import add_profiles

client = TestClient(app)

# Design 7.1 shape: directory sheets use the business name as their key
# (no synthetic id column), which is exactly the same-name/value-overlap
# scenario Day 11 link inference targets.
CUSTOMER_COLUMNS = ["客户名称", "区域", "联系电话"]
CUSTOMER_ROWS = [
    ["杭州云栖科技", "华东", "13800001111"],
    ["深圳前海贸易", "华南", "13900002222"],
    ["成都锦江商贸", "华东", "13700003333"],
    ["北京海淀贸易", "华北", "13600004444"],
    ["南京玄武商贸", "华东", "13500005555"],
    ["武汉江汉贸易", "华南", "13400006666"],
]
PRODUCT_COLUMNS = ["商品名称", "单价", "规格", "库存数量"]
PRODUCT_ROWS = [
    ["无线鼠标", "￥99.00", "黑色", "120"],
    ["机械键盘", "￥299.00", "黑色", "45"],
    ["显示器支架", "￥159.00", "银色", "8"],
    ["USB-C线材", "￥29.90", "白色", "300"],
    ["内胆包", "￥49.00", "灰色", "60"],
]
ORDER_COLUMNS = ["订单编号", "客户名称", "商品名称", "金额", "下单日期"]
ORDER_ROWS = [
    ["SO-1001", "杭州云栖科技", "无线鼠标", "99.0", "2026-07-03"],
    ["SO-1002", "杭州云栖科技", "机械键盘", "299.0", "2026-07-04"],
    ["SO-1003", "深圳前海贸易", "无线鼠标", "99.0", "2026-08-15"],
    ["SO-1004", "成都锦江商贸", "显示器支架", "159.0", "2026-09-01"],
    ["SO-1005", "北京海淀贸易", "USB-C线材", "29.9", "2026-09-20"],
    ["SO-1006", "南京玄武商贸", "内胆包", "49.0", "2026-10-02"],
    ["SO-1007", "武汉江汉贸易", "机械键盘", "299.0", "2026-10-05"],
    ["SO-1008", "杭州云栖科技", "无线鼠标", "99.0", "2026-10-06"],
]
CONFIG_ROWS = [["系统名称", "销售管理"], ["会话超时", "30"], ["总部城市", "上海"]]


def _sheet(
    name: str, columns: list[str], rows: Sequence[Sequence[object]], empty: bool = False
) -> ParsedSheet:
    if empty:
        return ParsedSheet(
            name=name, is_empty=True, header_row_number=None, columns=[],
            rows=[], data_row_count=0,
        )
    normalized = [[None if c is None else str(c) for c in row] for row in rows]
    return ParsedSheet(
        name=name, header_row_number=1, columns=columns,
        rows=normalized, data_row_count=len(normalized),
    )


def _run(sheets: list[ParsedSheet], file_name: str = "2026年销售台账.xlsx"):
    workbook = ParsedWorkbook(file_name=file_name, file_type="xlsx", sheets=sheets)
    return add_profiles(workbook)


def _yaml_model(result) -> BusinessModel:
    assert result.business_model_yaml is not None
    return model_from_yaml(result.business_model_yaml)


T = TypeVar("T")


def _view(model: BusinessModel, key: str, kind: type[T]) -> T:
    return next(
        v for v in model.views if isinstance(v, kind) and v.key == key
    )


def _ledger():
    return _run([
        _sheet("客户", CUSTOMER_COLUMNS, CUSTOMER_ROWS),
        _sheet("商品", PRODUCT_COLUMNS, PRODUCT_ROWS),
        _sheet("订单", ORDER_COLUMNS, ORDER_ROWS),
        _sheet("Sheet9", ["参数名", "参数值"], CONFIG_ROWS),
    ])


# ---------- link inference ----------


def test_customer_and_product_links_to_order_are_strong():
    result = _ledger()
    links = {link.key: link for link in result.inferred_links}
    assert set(links) == {"customer_to_order", "product_to_order"}

    customer_link = links["customer_to_order"]
    assert customer_link.from_entity == "customer"
    assert customer_link.to_entity == "order"
    assert customer_link.on_from == "customer_name"
    assert customer_link.on_to == "customer_name"
    assert customer_link.confidence == 0.88
    assert customer_link.needs_review is True  # every link is review-gated
    assert customer_link.overlap_recall == 1.0
    assert customer_link.overlap_precision == 1.0
    assert "value_set_overlap" in customer_link.signals
    assert customer_link.review_reason

    product_link = links["product_to_order"]
    assert product_link.on_from == "product_name"
    assert product_link.on_to == "product_name"
    assert product_link.confidence == 0.88


def test_same_name_non_identifier_columns_do_not_link():
    # Both sheets carry 备注 and 状态 with repeated values: generic columns
    # must never create name-only links.
    # Parent key column (客户编号) does not exist on the child sheet; the
    # only shared labels are generic non-identifier columns.
    a_rows = [["A1", "老客户", "1"], ["A2", "新客户", "0"], ["A3", "老客户", "1"]]
    b_rows = [["老客户", "1"], ["老客户", "0"], ["新客户", "1"]]
    result = _run([
        _sheet("客户", ["客户编号", "备注", "状态"], a_rows),
        _sheet("订单", ["备注", "状态", "金额", "下单日期"],
               [r + ["100", "2026-01-01"] for r in b_rows]),
    ])
    assert result.inferred_links == []


def test_same_name_without_value_overlap_is_low_confidence():
    # Parent directory keyed by 名称; child repeats a same-name column whose
    # values never match the parent set.
    parent_rows = [["张伟"], ["李娜"], ["王强"]]
    child_rows = [
        ["SO-1", "赵敏", "100", "2026-01-01"],
        ["SO-2", "周杰", "200", "2026-02-01"],
        ["SO-3", "赵敏", "300", "2026-03-01"],
    ]
    result = _run([
        _sheet("客户", ["名称"], parent_rows),
        _sheet("订单", ["订单编号", "名称", "金额", "下单日期"], child_rows),
    ])
    keys = {link.key: link for link in result.inferred_links}
    assert "customer_to_order" in keys
    weak = keys["customer_to_order"]
    assert weak.confidence == 0.63 and weak.needs_review is True
    assert weak.overlap_recall == 0.0


def test_unique_both_sides_is_not_one_to_many():
    parent_rows = [["张伟"], ["李娜"]]
    # child 名称 also unique: one-to-one (or coincidence), not supported.
    child_rows = [
        ["SO-1", "王强", "100", "2026-01-01"],
        ["SO-2", "周杰", "200", "2026-02-01"],
    ]
    result = _run([
        _sheet("客户", ["名称"], parent_rows),
        _sheet("订单", ["订单编号", "名称", "金额", "下单日期"], child_rows),
    ])
    assert result.inferred_links == []


def test_different_name_but_same_values_links_when_semantic():
    # Parent key 名称, child column 客户名称 (different label, same values).
    parent_rows = [["杭州云栖"], ["深圳前海"], ["成都锦江"]]
    child_rows = [
        ["SO-1", "杭州云栖", "100", "2026-01-01"],
        ["SO-2", "深圳前海", "200", "2026-02-01"],
        ["SO-3", "成都锦江", "300", "2026-03-01"],
        ["SO-4", "杭州云栖", "400", "2026-04-01"],
    ]
    result = _run([
        _sheet("客户", ["名称"], parent_rows),
        _sheet("订单", ["订单编号", "客户名称", "金额", "下单日期"], child_rows),
    ])
    keys = {link.key: link for link in result.inferred_links}
    link = keys.get("customer_to_order")
    assert link is not None and link.confidence == 0.75
    assert link.on_to == "customer_name"


# ---------- assembly ----------


def test_model_is_valid_with_three_entities_and_chinese_fields_mapped():
    result = _ledger()
    assert result.business_model is not None
    assert result.business_model_yaml

    model = _yaml_model(result)  # domain validation
    assert isinstance(model, BusinessModel)
    assert [e.key for e in model.entities] == ["customer", "product", "order"]

    order = next(e for e in model.entities if e.key == "order")
    assert order.key_field == "order_no"
    field_map = {f.key: f for f in order.fields}
    assert {"order_no", "customer_name", "product_name", "amount", "order_date"} <= set(field_map)
    assert field_map["amount"].type.value == "money"
    assert field_map["amount"].role.value == "measure"
    # enum candidate values survive into the domain field
    customer = next(e for e in model.entities if e.key == "customer")
    region = next(f for f in customer.fields if f.key == "region")
    assert region.type.value == "enum" and region.values == ["华东", "华北", "华南"]


def test_metric_only_for_flow_money_and_views_reference_it():
    result = _ledger()
    model = _yaml_model(result)

    metric_keys = {m.key for m in model.metrics}
    # 单价 exists on product but is a price, not summed.
    assert metric_keys == {"total_amount"}
    metric = model.metrics[0]
    assert metric.entity == "order"
    assert metric.formula.op.value == "sum" and metric.formula.field == "amount"
    assert metric.business_definition

    dashboard = _view(model, "home_dashboard", DashboardView)
    assert dashboard.metrics == ["total_amount"]

    order_list = _view(model, "order_list", ListView)
    assert order_list.columns[0] == "order_no"
    assert "amount" in order_list.columns and "order_date" in order_list.columns
    assert order_list.sorts[0].field == "order_date"
    assert order_list.sorts[0].dir.value == "desc"
    assert "customer_name" in order_list.filters

    customer_detail = _view(model, "customer_detail", DetailView)
    related = [b for b in customer_detail.blocks if b.type.value == "related_list"]
    assert len(related) == 1 and related[0].link == "customer_to_order"

    assert [item.view for item in model.navigation] == [
        "home_dashboard", "customer_list", "product_list", "order_list",
    ]


def test_unknown_and_keyless_sheets_are_skipped_with_notes():
    # config sheet is unknown; an empty sheet is ignored.
    result = _run([
        _sheet("Sheet9", ["参数名", "参数值"], CONFIG_ROWS),
        _sheet("空表", [], [], empty=True),
    ])
    assert result.business_model is None
    assert result.inferred_links == []
    assert any("unknown" in note for note in result.assembly_notes)


def test_single_sheet_order_assembles_without_links():
    result = _run([_sheet("订单", ORDER_COLUMNS, ORDER_ROWS)])
    model = _yaml_model(result)
    assert len(model.entities) == 1 and model.entities[0].key == "order"
    assert model.links == []
    detail = _view(model, "order_detail", DetailView)
    assert [b.type.value for b in detail.blocks] == ["fields"]
    assert result.assembly_notes == []


def test_yaml_round_trip_and_alias_handling():
    result = _ledger()
    yaml_text = result.business_model_yaml
    assert yaml_text is not None
    # link alias `from` must survive the YAML 1.2 loader
    assert "from: customer" in yaml_text
    model = model_from_yaml(yaml_text)
    link = next(lk for lk in model.links if lk.key == "customer_to_order")
    assert link.from_ == "customer" and link.to == "order"
    assert link.on.from_ == "customer_name"


def test_duplicate_entity_kind_gets_namespaced_key():
    # Two order-shaped sheets -> order and order_2, both valid entities.
    second_rows = [
        ["R1", "9", "2026-01-01"], ["R2", "9", "2026-02-01"],
        ["R3", "8", "2026-03-01"],
    ]
    result = _run([
        _sheet("订单", ORDER_COLUMNS, ORDER_ROWS),
        _sheet("退货", ["退货编号", "金额", "退货日期"], second_rows),
    ])
    model = _yaml_model(result)
    assert [e.key for e in model.entities] == ["order", "order_2"]
    views = {v.key for v in model.views}
    assert "order_list" in views and "order_2_list" in views


# ---------- end to end ----------


def test_parse_endpoint_returns_valid_model_from_real_xlsx():
    wb = openpyxl.Workbook()
    wb.remove(wb.worksheets[0])
    for name, columns, rows in (
        ("客户", CUSTOMER_COLUMNS, CUSTOMER_ROWS),
        ("商品", PRODUCT_COLUMNS, PRODUCT_ROWS),
        ("订单", ORDER_COLUMNS, ORDER_ROWS),
        ("Sheet9", ["参数名", "参数值"], CONFIG_ROWS),
    ):
        ws = wb.create_sheet(name)
        ws.append(columns)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)

    response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("2026年销售台账.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["business_model"] is not None
    assert payload["business_model_yaml"]
    keys = {link["key"] for link in payload["inferred_links"]}
    assert keys == {"customer_to_order", "product_to_order"}
    # the returned model re-validates through the domain schema
    model = model_from_yaml(payload["business_model_yaml"])
    assert model.app.source is not None
    assert model.app.source.files == ["2026年销售台账.xlsx"]
    assert model.app.name == "2026年销售台账"
