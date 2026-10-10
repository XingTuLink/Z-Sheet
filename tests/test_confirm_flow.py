"""Tests for Day 15: Confirm 全链路 (review decisions -> model + rows -> app).

The tests exercise the HTTP contract end to end:

  parse/stream (collect every needs_review item like the UI does)
    -> /ingestion/confirm
    -> /runtime/workspace/bootstrap serving the user's own rows
"""

from __future__ import annotations

import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.understanding.confirmation import (
    ConfirmDecisions,
    ConfirmError,
    LinkDecisionPayload,
    apply_link_decisions,
    missing_link_decisions,
)
from backend.understanding.schemas import InferredLink

client = TestClient(app)


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    customers = wb.worksheets[0]
    customers.title = "客户"
    customers.append(["客户名称", "区域", "联系电话"])
    customers.append(["杭州云栖", "华东", "13800000001"])
    customers.append(["深圳前海", "华南", "13800000002"])
    customers.append(["北京中关村", "华北", "13800000003"])
    customers.append(["广州天河", "华南", "13800000004"])
    customers.append(["成都高新", "华东", "13800000005"])
    customers.append(["武汉光谷", "华北", "13800000006"])

    products = wb.create_sheet("商品")
    products.append(["商品名称", "单价"])
    products.append(["云主机", 199])
    products.append(["对象存储", 49])
    products.append(["带宽包", 89])

    orders = wb.create_sheet("订单")
    orders.append(["订单编号", "客户名称", "商品名称", "金额", "下单日期"])
    rows = [
        ("SO-1", "杭州云栖", "云主机", 199, "2026/1/3"),
        ("SO-2", "深圳前海", "云主机", 199, "2026/1/5"),
        ("SO-3", "北京中关村", "带宽包", 89, "2026/1/8"),
        ("SO-4", "广州天河", "对象存储", 49, "2026/1/9"),
        ("SO-5", "成都高新", "云主机", 199, "2026/2/1"),
        ("SO-6", "武汉光谷", "带宽包", 89, "2026/2/2"),
        ("SO-7", "杭州云栖", "对象存储", 49, "2026/2/3"),
        ("SO-8", "深圳前海", "云主机", 199, "2026/2/4"),
    ]
    for row in rows:
        orders.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_result() -> dict:
    resp = client.post(
        "/api/v1/ingestion/parse/stream",
        files={
            "file": (
                "台账.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200
    for block in resp.text.strip().split("\n\n"):
        if block.startswith("event: result"):
            return json.loads(block.split("data: ", 1)[1])
    raise AssertionError("no result frame")


def _acknowledged(result: dict) -> list[str]:
    """Collect every needs_review entity/field/metric id exactly like the UI."""
    model = result["business_model"]
    sheet_by_name = {sheet["name"]: sheet for sheet in result["sheets"]}
    ids: list[str] = []
    for entity in model["entities"]:
        sheet_name = (entity.get("source") or {}).get("sheet")
        inferred = sheet_by_name.get(sheet_name, {}).get("inferred_entity")
        if inferred and inferred["needs_review"]:
            ids.append(entity["key"])
        for field in entity["fields"]:
            if field["needs_review"]:
                ids.append(f"{entity['key']}.{field['key']}")
    for metric in model.get("metrics", []):
        if metric["needs_review"]:
            ids.append(f"metric:{metric['key']}")
    return ids


def _decisions(
    result: dict, *, reject: set[str] | None = None
) -> dict:
    reject = reject or set()
    links: dict[str, dict[str, str]] = {}
    for link in result["inferred_links"]:
        if link["key"] in reject:
            links[link["key"]] = {"decision": "rejected"}
        else:
            links[link["key"]] = {
                "decision": "accepted",
                "on_from": link["on_from"],
                "on_to": link["on_to"],
            }
    return {"acknowledged": _acknowledged(result), "links": links}


def _confirm(result: dict, *, reject: set[str] | None = None):
    return client.post(
        "/api/v1/ingestion/confirm",
        files={
            "file": (
                "台账.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"decisions": json.dumps(_decisions(result, reject=reject))},
    )


def test_confirm_persists_model_and_workbook_rows():
    result = _parse_result()
    assert {link["key"] for link in result["inferred_links"]} == {
        "customer_to_order",
        "product_to_order",
    }

    resp = _confirm(result)
    assert resp.status_code == 200, resp.text
    summary = resp.json()
    assert summary["app_key"] == "workspace"
    assert summary["version"] == 1
    assert set(summary["links_accepted"]) == {
        "customer_to_order",
        "product_to_order",
    }
    assert summary["links_rejected"] == []
    assert summary["record_counts"] == {"customer": 6, "product": 3, "order": 8}

    boot = client.get("/api/v1/runtime/workspace/bootstrap")
    assert boot.status_code == 200
    payload = boot.json()
    assert payload["app_key"] == "workspace"
    assert payload["version"] == 1
    assert {e["key"] for e in payload["model"]["entities"]} == {
        "customer",
        "product",
        "order",
    }
    # The generated app renders the user's rows, not the demo seed.
    assert len(payload["records"]["customer"]) == 6
    assert len(payload["records"]["product"]) == 3
    orders = payload["records"]["order"]
    assert len(orders) == 8
    assert orders[0]["order_no"] == "SO-1"
    assert isinstance(orders[0]["amount"], (int, float))
    # "2026/1/3" style cells are normalized to ISO dates.
    assert orders[0]["order_date"] == "2026-01-03"
    # Both links survived and power detail related-list blocks.
    assert {link["key"] for link in payload["model"]["links"]} == {
        "customer_to_order",
        "product_to_order",
    }


def test_confirm_blocked_until_review_queue_cleared():
    resp = client.post(
        "/api/v1/ingestion/confirm",
        files={
            "file": (
                "台账.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"decisions": json.dumps({"acknowledged": [], "links": {}})},
    )
    assert resp.status_code == 422
    body = resp.json()
    unresolved = body["unresolved"]
    assert "relation:customer_to_order" in unresolved
    assert "relation:product_to_order" in unresolved
    # Nothing was persisted.
    assert client.get("/api/v1/runtime/workspace/bootstrap").status_code == 404


def test_confirm_reject_link_rebuilds_views_without_it():
    result = _parse_result()
    resp = _confirm(result, reject={"product_to_order"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["links_rejected"] == ["product_to_order"]

    payload = client.get("/api/v1/runtime/workspace/bootstrap").json()
    assert [link["key"] for link in payload["model"]["links"]] == [
        "customer_to_order"
    ]
    # No related-list block may reference the dropped link (domain
    # validation already rejects dangling references, and assembly rebuilds
    # views from the surviving links).
    dumped = json.dumps(payload["model"], ensure_ascii=False)
    assert "product_to_order" not in dumped
    # Order rows are still persisted; only the navigation relationship is gone.
    assert len(payload["records"]["order"]) == 8


def test_confirm_twice_appends_a_new_version():
    result = _parse_result()
    assert _confirm(result).status_code == 200
    second = _confirm(result, reject={"product_to_order"})
    assert second.status_code == 200
    assert second.json()["version"] == 2

    payload = client.get("/api/v1/runtime/workspace/bootstrap").json()
    assert payload["version"] == 2
    assert [link["key"] for link in payload["model"]["links"]] == [
        "customer_to_order"
    ]
    assert len(payload["records"]["order"]) == 8


def test_demo_bootstrap_stays_independent():
    result = _parse_result()
    assert _confirm(result).status_code == 200

    demo = client.get("/api/v1/runtime/default/bootstrap")
    assert demo.status_code == 200
    payload = demo.json()
    assert payload["app_key"] == "default"
    # Demo seed rows still come from the checked-in fixtures.
    assert payload["records"]
    assert "customer" in payload["records"]
    assert payload["model"]["app"]["name"] != "台账"


def test_confirm_rejects_malformed_decisions_json():
    resp = client.post(
        "/api/v1/ingestion/confirm",
        files={
            "file": (
                "台账.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"decisions": "{not json"},
    )
    assert resp.status_code == 422


def test_apply_link_decisions_unit():
    links = [
        InferredLink(
            key="customer_to_order",
            from_entity="customer",
            to_entity="order",
            on_from="name",
            on_to="customer_name",
            confidence=0.9,
            needs_review=True,
            overlap_recall=1.0,
            overlap_precision=1.0,
        )
    ]
    decisions = ConfirmDecisions(
        links={"customer_to_order": LinkDecisionPayload(decision="accepted")}
    )
    kept, accepted, rejected = apply_link_decisions(links, decisions.links)
    assert rejected == []
    assert accepted == ["customer_to_order"]
    assert kept[0].on_from == "name"  # None override keeps the inference

    override = ConfirmDecisions(
        links={
            "customer_to_order": LinkDecisionPayload(
                decision="accepted", on_from="code", on_to="buyer"
            )
        }
    )
    kept, _, _ = apply_link_decisions(links, override.links)
    assert (kept[0].on_from, kept[0].on_to) == ("code", "buyer")

    dropped = ConfirmDecisions(
        links={"customer_to_order": LinkDecisionPayload(decision="rejected")}
    )
    kept, _, rejected = apply_link_decisions(links, dropped.links)
    assert kept == [] and rejected == ["customer_to_order"]

    assert missing_link_decisions(links, {}) == ["relation:customer_to_order"]

    with pytest.raises(ConfirmError):
        apply_link_decisions(
            links,
            {
                "customer_to_order": LinkDecisionPayload(decision="accepted"),
                "ghost": LinkDecisionPayload(decision="rejected"),
            },
        )
