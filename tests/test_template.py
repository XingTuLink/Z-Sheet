"""Tests for the downloadable standard workbook template (Day 15 follow-up)."""

from __future__ import annotations

import io

import openpyxl
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.domain.models import BusinessModel
from backend.ingestion.parser import parse_workbook
from backend.ingestion.template import TEMPLATE_FILENAME, build_template_workbook
from backend.understanding.profiling import add_profiles

client = TestClient(app)


def test_download_template_endpoint():
    resp = client.get("/api/v1/ingestion/template")
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"  # xlsx is a zip container
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert "filename*=UTF-8''" in disposition
    assert TEMPLATE_FILENAME.endswith(".xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["客户", "商品", "订单"]
    assert [cell.value for cell in wb["客户"][1]] == ["客户名称", "区域", "联系电话"]
    assert [cell.value for cell in wb["商品"][1]] == ["商品名称", "单价"]
    assert [cell.value for cell in wb["订单"][1]] == [
        "订单编号", "客户名称", "商品名称", "金额", "下单日期",
    ]
    # Example data ships with the template; header row is frozen.
    assert wb["客户"].freeze_panes == "A2"
    assert wb["客户"].max_row == 7
    assert wb["订单"].max_row == 9


def test_template_itself_assembles_into_recognizable_app():
    """The template must pass the pipeline it teaches — executable docs."""
    content = build_template_workbook().getvalue()
    result = add_profiles(parse_workbook("template.xlsx", content))

    assert result.business_model is not None
    model = BusinessModel.model_validate(result.business_model)
    assert {entity.key for entity in model.entities} == {
        "customer", "product", "order",
    }
    assert {link.key for link in result.inferred_links} == {
        "customer_to_order", "product_to_order",
    }
    assert result.assembly_notes == []
