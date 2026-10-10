"""Day 22: punctuation-row hardening + server-side understanding sessions.

Confirm must consume exactly the understanding result the user reviewed:
parse stores it keyed by file hash, confirm loads it instead of re-inferring
(which with an LLM would be non-deterministic and could disagree).
"""

from __future__ import annotations

import io
import json
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.ai.provider import ChatClient, LLMConfig
from backend.api.main import app
from backend.ingestion.parser import parse_workbook
from backend.ingestion.schemas import ParsedSheet, ParsedWorkbook
from backend.storage.db import SessionLocal
from backend.storage.repositories import (
    model_repository,
    understanding_repository,
)
from backend.understanding import confirmation
from backend.understanding.confirmation import (
    WORKSPACE_APP_KEY,
    ConfirmDecisions,
    confirm_workbook,
    unresolved_review_items,
)
from backend.understanding.model_assembly import assemble_model, build_entity_plans
from backend.understanding.profiling import add_profiles


def _workbook(sheets: list[ParsedSheet]):
    return ParsedWorkbook(file_name="t.xlsx", file_type="xlsx", sheets=sheets)


client = TestClient(app)

LLM_CONFIG = LLMConfig(
    base_url="https://api.deepseek.com",
    api_key="sk-test-1234567890",
    model="deepseek-flash",
)


class ProposalTransport:
    def __init__(self, proposal: dict[str, Any]) -> None:
        self.proposal = proposal

    def __call__(
        self,
        url: str,
        headers: dict[str, str],
        payload: dict[str, Any],
        timeout: float,
    ) -> dict[str, Any]:
        return {
            "choices": [
                {"message": {"content": json.dumps(self.proposal, ensure_ascii=False)}}
            ]
        }


def _review_xlsx() -> bytes:
    """The external-review pattern: 1 real row + a punctuation pseudo-row."""
    wb = openpyxl.Workbook()
    sheet = wb.worksheets[0]
    sheet.title = "校对总表"
    sheet.append(["序号", "姓名", "性别", "联系电话", "备注"])
    sheet.append(["1", "郭婉莹", "女", "13759943324", "ok"])
    # Artifact row: only a "." in the last column.
    sheet.append([None, None, None, None, "."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _llm_proposal() -> dict[str, Any]:
    return {
        "entities": [
            {
                "source_sheet": "校对总表",
                "key": "proofreader",
                "name": "外审人员",
                "key_field": "姓名",
                "confidence": 0.95,
                "needs_review": False,
                "reason": "按人登记",
                "fields": [
                    {"column": "姓名", "type": "string", "role": "identifier",
                     "confidence": 0.95},
                    {"column": "性别", "type": "enum", "role": "enum",
                     "confidence": 0.95, "enum_values": ["男", "女"]},
                    {"column": "联系电话", "type": "phone", "role": "text",
                     "confidence": 0.95},
                    {"column": "备注", "type": "string", "role": "text",
                     "confidence": 0.9},
                    {"column": "序号", "type": "number", "role": "dimension",
                     "confidence": 0.8},
                ],
            }
        ],
        "links": [],
    }


def _review_decisions(result) -> dict[str, Any]:
    """Every needs-review item acknowledged (bare ids), every link decided."""
    plans, _ = build_entity_plans(result)
    model = assemble_model(result, plans, result.inferred_links)
    assert model is not None
    # The decisions payload uses bare ids; unresolved_review_items reports the
    # same items with "entity:"/"field:" display prefixes.
    prefixed = unresolved_review_items(
        plans, model, ConfirmDecisions.model_validate({})
    )
    acknowledged = [item.split(":", 1)[1] for item in prefixed]
    links = {
        link.key: {"decision": "accepted", "on_from": link.on_from,
                   "on_to": link.on_to}
        for link in result.inferred_links
    }
    return {"acknowledged": acknowledged, "links": links}


def test_parser_drops_punctuation_only_rows() -> None:
    workbook = parse_workbook(
        "t.csv",
        "序号,姓名,备注\n1,郭婉莹,ok\n,.,.\n2,李明,-\n-,-,/\n".encode(),
    )
    sheet = workbook.sheets[0]
    assert sheet.data_row_count == 2
    assert [row[1] for row in sheet.rows] == ["郭婉莹", "李明"]


def test_real_external_review_row_survives_in_xlsx() -> None:
    workbook = parse_workbook("校对.xlsx", _review_xlsx())
    sheet = workbook.sheets[0]
    assert sheet.data_row_count == 1
    assert sheet.rows[0][1] == "郭婉莹"


def test_parser_drops_aggregate_total_rows() -> None:
    workbook = parse_workbook(
        "报价.csv",
        "序号,功能,金额\n1,首页,150\n2,详情,100\n,,合计\n".encode(),
    )
    sheet = workbook.sheets[0]
    assert sheet.data_row_count == 2
    assert [row[1] for row in sheet.rows] == ["首页", "详情"]
    assert any("aggregate" in warning for warning in sheet.warnings)


def test_parser_drops_total_row_in_any_column_xlsx() -> None:
    wb = openpyxl.Workbook()
    sheet = wb.worksheets[0]
    sheet.title = "报价单"
    sheet.append(["序号", "模块", "功能", "描述", "数量", "单价"])
    sheet.append(["1", "小程序", "首页", "展示", "1", "150"])
    # Real-world shape: the marker sits in a non-key column, totals on right.
    sheet.append([None, None, None, "合计", "3138", None])
    buf = io.BytesIO()
    wb.save(buf)
    workbook = parse_workbook("小程序报价单.xlsx", buf.getvalue())
    parsed_sheet = workbook.sheets[0]
    assert parsed_sheet.data_row_count == 1
    assert parsed_sheet.rows[0][2] == "首页"


def test_understanding_session_round_trips_and_overwrites() -> None:
    workbook = parse_workbook("校对.xlsx", _review_xlsx())
    result = add_profiles(workbook, llm_client=None)
    digest = understanding_repository.session_key(_review_xlsx())

    db = SessionLocal()
    try:
        assert understanding_repository.load_understanding(db, digest) is None
        understanding_repository.save_understanding(db, digest, result)
        loaded = understanding_repository.load_understanding(db, digest)
        assert loaded is not None
        assert loaded.understanding_engine == "rules"
        assert loaded.sheets[0].columns == result.sheets[0].columns

        # Same file re-parsed: overwrite, not duplicate.
        understanding_repository.save_understanding(db, digest, result)
        again = understanding_repository.load_understanding(db, digest)
        assert again is not None
    finally:
        db.close()


def test_confirm_uses_cached_llm_snapshot_without_re_inferring() -> None:
    content = _review_xlsx()
    workbook = parse_workbook("校对.xlsx", content)
    # Parse-time understanding used the LLM (fake, offline).
    llm_result = add_profiles(
        workbook,
        llm_client=ChatClient(LLM_CONFIG, transport=ProposalTransport(_llm_proposal())),
    )
    assert llm_result.understanding_engine == "llm"

    db = SessionLocal()
    try:
        understanding_repository.save_understanding(
            db, understanding_repository.session_key(content), llm_result
        )
        # No LLM is configured in the test environment: a re-run would yield
        # the rules engine, whose entity key differs from "proofreader".
        decisions = _review_decisions(llm_result)
        summary = confirm_workbook(db, "校对.xlsx", content, decisions)
        record = model_repository.get_current(db, summary.app_key)
        keys = {entity["key"] for entity in record.snapshot["entities"]}
        assert "proofreader" in keys
        assert summary.record_counts == {"proofreader": 1}
    finally:
        db.close()


def test_confirm_falls_back_to_pipeline_on_cache_miss() -> None:
    content = _review_xlsx()
    # Nothing stored for this hash: confirm re-runs the (rules) pipeline and
    # still succeeds, with the punctuation row filtered at parse time.
    db = SessionLocal()
    try:
        workbook = parse_workbook("校对.xlsx", content)
        rules_result = add_profiles(workbook, llm_client=None)
        decisions = _review_decisions(rules_result)
        summary = confirm_workbook(db, "校对.xlsx", content, decisions)
        assert summary.app_key == WORKSPACE_APP_KEY
        assert summary.version == 1
        # Exactly one real row (the "." pseudo-row never becomes data).
        assert sum(summary.record_counts.values()) == 1
    finally:
        db.close()


def test_parse_endpoint_persists_session_for_confirm() -> None:
    content = _review_xlsx()
    resp = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("校对.xlsx", content,
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    parsed = resp.json()

    db = SessionLocal()
    try:
        stored = understanding_repository.load_understanding(
            db, understanding_repository.session_key(content)
        )
        assert stored is not None
        assert stored.understanding_engine == parsed["understanding_engine"]
    finally:
        db.close()


def test_blank_key_row_at_confirm_raises_confirm_error_not_value_error() -> None:
    """A data row missing its key must be a reviewable ConfirmError (422)."""
    # 39 fully-populated rows pass the 95% key gate; row 40 lacks the key but
    # carries dimension data, so it is neither blank nor an aggregate row.
    rows: list[list[str | None]] = [
        [f"客户{i}", "华东"] for i in range(39)
    ]
    rows.append([None, "华北"])
    sheet = ParsedSheet(
        name="客户",
        header_row_number=1,
        columns=["客户名称", "区域"],
        rows=rows,
        data_row_count=40,
    )
    proposal = {
        "entities": [
            {"source_sheet": "客户", "key": "customer", "name": "客户",
             "key_field": "客户名称", "confidence": 0.95,
             "fields": [
                 {"column": "客户名称", "type": "string", "role": "identifier",
                  "confidence": 0.95},
                 {"column": "区域", "type": "string", "role": "dimension",
                  "confidence": 0.9},
             ]}
        ],
        "links": [],
    }
    result = add_profiles(
        _workbook([sheet]),
        llm_client=ChatClient(LLM_CONFIG, transport=ProposalTransport(proposal)),
    )
    plans, _ = build_entity_plans(result)
    model = assemble_model(result, plans, [])
    assert model is not None and len(plans) == 1

    with pytest.raises(confirmation.ConfirmError, match="is required|主键|校验"):
        confirmation.extract_records(plans, model)
