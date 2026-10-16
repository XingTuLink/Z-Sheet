"""Tests for Day 5 ingestion: XLSX/CSV parsing, sheet discovery, header detection."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.api.main import app
from backend.api.routes import ingestion as ingestion_route
from backend.ingestion.parser import (
    MAX_COLUMNS,
    MAX_DATA_ROWS,
    ParseError,
    parse_workbook,
)

client = TestClient(app)


def make_xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- XLSX ----------


def test_xlsx_discovers_multiple_sheets_including_empty_one():
    content = make_xlsx_bytes(
        {
            "客户": [["客户名", "区域"], ["张三", "华东"], ["李四", "西南"]],
            "订单": [["订单号", "金额"], ["SO-1", 100], ["SO-2", 200]],
            "说明": [],
        }
    )
    workbook = parse_workbook("sales.xlsx", content)

    assert workbook.file_type == "xlsx"
    assert [sheet.name for sheet in workbook.sheets] == ["客户", "订单", "说明"]

    customer, order, notes = workbook.sheets
    assert customer.header_row_number == 1
    assert customer.columns == ["客户名", "区域"]
    assert customer.rows == [["张三", "华东"], ["李四", "西南"]]
    assert customer.data_row_count == 2
    assert order.rows[0] == ["SO-1", "100"]
    assert notes.is_empty is True
    assert notes.header_row_number is None


def test_xlsx_header_below_banner_and_blank_rows():
    content = make_xlsx_bytes(
        {
            "订单明细": [
                ["2026 年销售明细表"],   # banner: single cell, low coverage
                [],                      # blank gap
                ["订单号", "客户", "金额", "日期"],
                # Excel has no pure date type: a written date round-trips as
                # datetime via openpyxl values_only, so expect ISO datetimes.
                ["SO-1", "张三", 87350.5, dt.datetime(2026, 10, 7)],
                [],                      # blank row inside data must be dropped
                ["SO-2", "李四", 99, dt.datetime(2026, 10, 8)],
            ]
        }
    )
    sheet = parse_workbook("sales.xlsx", content).sheets[0]

    assert sheet.header_row_number == 3
    assert sheet.columns == ["订单号", "客户", "金额", "日期"]
    assert sheet.rows == [
        ["SO-1", "张三", "87350.5", "2026-10-07T00:00:00"],
        ["SO-2", "李四", "99", "2026-10-08T00:00:00"],
    ]


def test_xlsx_native_cell_types_become_strings_or_none():
    content = make_xlsx_bytes(
        {
            "类型": [
                ["a", "b", "c", "d", "e", "f"],
                [1, 1.5, True, dt.datetime(2026, 1, 2, 3, 4, 5), None, "  spaced  "],
            ]
        }
    )
    row = parse_workbook("types.xlsx", content).sheets[0].rows[0]
    assert row == ["1", "1.5", "TRUE", "2026-01-02T03:04:05", None, "spaced"]


def test_xlsx_header_only_sheet_warns():
    content = make_xlsx_bytes({"空表": [["列1", "列2"]]})
    sheet = parse_workbook("h.xlsx", content).sheets[0]
    assert sheet.data_row_count == 0
    assert any("no data rows" in warning for warning in sheet.warnings)


def test_xlsx_vertical_merged_category_is_forward_filled():
    # Real-world layout (栈上月明业务清单): a category cell is vertically
    # merged across the service rows belonging to it; only the anchor cell
    # carries the value. All covered rows must inherit it.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "业务清单"
    ws.append(["序号", "业务板块", "服务项目"])
    ws.append([1, "板块甲", "客服"])
    ws.append([2, None, "跟单"])
    ws.append([3, None, "报表"])
    ws.append([4, "板块乙", "官网"])
    ws.merge_cells("B2:B4")
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("services.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.header_row_number == 1
    assert sheet.rows == [
        ["1", "板块甲", "客服"],
        ["2", "板块甲", "跟单"],
        ["3", "板块甲", "报表"],
        ["4", "板块乙", "官网"],
    ]


def test_xlsx_merged_anchor_line_breaks_are_collapsed():
    # The anchor of a vertical category merge visually wraps with an
    # in-cell newline ("AI 智能体\n开发"); both the anchor row and the
    # filled rows must carry the same single-line enum value.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "业务清单"
    ws.append(["序号", "业务板块", "服务项目"])
    ws.append([1, "AI 智能体\n开发", "客服"])
    ws.append([2, None, "跟单"])
    ws.append([3, "普通板块\n保留换行", "推荐"])
    ws.merge_cells("B2:B3")  # vertical merge on the wrapped anchor only
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("services.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.rows[0][:2] == ["1", "AI 智能体开发"]
    assert sheet.rows[1][:2] == ["2", "AI 智能体开发"]
    # Non-merged cells keep their interior newlines untouched.
    assert sheet.rows[2][1] == "普通板块\n保留换行"


def test_xlsx_horizontal_merged_banner_is_not_filled_into_data():
    # Title/annotation rows are horizontally merged; they must NOT fan out
    # into other columns and must not become data rows.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "清单"
    ws.append(["公司宣传大标题", None, None])
    ws.append(["序号", "板块", "服务"])
    ws.append([1, "板块甲", "客服"])
    ws.merge_cells("A1:C1")
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("banner.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.header_row_number == 2
    assert sheet.columns == ["序号", "板块", "服务"]
    assert sheet.rows == [["1", "板块甲", "客服"]]
    assert sheet.data_row_count == 1


def test_xlsx_unnamed_and_duplicate_headers_warn():
    content = make_xlsx_bytes(
        {
            "表": [
                ["名称", None, "名称"],
                ["a", "b", "c"],
            ]
        }
    )
    sheet = parse_workbook("types.xlsx", content).sheets[0]
    assert sheet.columns == ["名称", "column_2", "名称"]
    assert any("unnamed" in warning for warning in sheet.warnings)
    assert any("duplicate" in warning for warning in sheet.warnings)


def test_xlsx_two_row_grouped_header_prefixes_duplicate_children():
    # Real layout: two groups share a repeated child name; the group label
    # disambiguates 收款信息_金额 / 退款信息_金额 instead of two bare 金额.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "对账"
    ws.append(["收款信息", None, "退款信息", None])
    ws.append(["日期", "金额", "日期", "金额"])
    ws.append(["2026-01-01", 100, "2026-01-02", 50])
    ws.merge_cells("A1:B1")
    ws.merge_cells("C1:D1")
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("recon.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.header_row_number == 2
    assert sheet.columns == [
        "收款信息_日期",
        "收款信息_金额",
        "退款信息_日期",
        "退款信息_金额",
    ]
    assert sheet.rows == [["2026-01-01", "100", "2026-01-02", "50"]]
    assert not any("duplicate" in warning for warning in sheet.warnings)


def test_xlsx_two_row_header_single_group_is_not_prefixed():
    # One merged group spanning all columns is a banner, not a group row;
    # child names stay unchanged.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "表"
    ws.append(["2026年度销售", None, None])
    ws.append(["名称", "数量", "金额"])
    ws.append(["甲", 1, 10])
    ws.merge_cells("A1:C1")
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("single.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.header_row_number == 2
    assert sheet.columns == ["名称", "数量", "金额"]


def test_xlsx_formula_cells_without_cached_value_warn_and_read_empty():
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "订单"
    ws.append(["名称", "数量", "单价", "金额"])
    ws.append(["甲", 2, 10, "=B2*C2"])
    ws.append(["乙", 3, 5, "=B3*C3"])
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("fx.xlsx", buffer.getvalue()).sheets[0]
    assert [row[3] for row in sheet.rows] == [None, None]
    assert any("formula" in warning for warning in sheet.warnings)


def test_xlsx_fullwidth_header_is_nfkc_normalized():
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "表"
    ws.append(["客户名称", "ＡＢＣ", "金额"])
    ws.append(["张三", "x", 100])
    buffer = io.BytesIO()
    wb.save(buffer)

    sheet = parse_workbook("fw.xlsx", buffer.getvalue()).sheets[0]
    assert sheet.columns == ["客户名称", "ABC", "金额"]


def test_ragged_rows_align_to_header_width():
    content = make_xlsx_bytes({"表": [["a", "b", "c"], ["x", "y"]]})
    sheet = parse_workbook("r.xlsx", content).sheets[0]
    assert sheet.rows == [["x", "y", None]]


def test_too_many_columns_rejected():
    content = make_xlsx_bytes({"宽表": [["h"] * (MAX_COLUMNS + 1), ["v"] * (MAX_COLUMNS + 1)]})
    with pytest.raises(ParseError, match="columns"):
        parse_workbook("wide.xlsx", content)


# ---------- CSV ----------


def test_csv_utf8_plain_and_bom():
    plain = "客户名,区域\n张三,华东\n".encode()
    sheet = parse_workbook("a.csv", plain).sheets[0]
    assert sheet.columns == ["客户名", "区域"]
    assert sheet.rows == [["张三", "华东"]]

    bom = "客户名,区域\n张三,华东\n".encode("utf-8-sig")
    sheet_bom = parse_workbook("b.csv", bom).sheets[0]
    assert sheet_bom.columns == ["客户名", "区域"]


def test_csv_gbk_encoding_from_excel_cn():
    content = "客户名,区域\n张三,华东\n李四,西南\n".encode("gbk")
    workbook = parse_workbook("gbk.csv", content)
    sheet = workbook.sheets[0]
    assert workbook.file_type == "csv"
    assert sheet.name == "gbk"
    assert sheet.rows == [["张三", "华东"], ["李四", "西南"]]


def test_csv_utf16_bom_is_decoded_and_gbk_is_not_mistaken_for_utf16():
    utf16 = parse_workbook("u.csv", "a,b\n1,2\n".encode("utf-16")).sheets[0]
    assert utf16.columns == ["a", "b"]
    assert utf16.rows == [["1", "2"]]

    # Regression: GBK byte streams are even length and native-endian utf-16
    # would "decode" them into Korean-looking garbage without a BOM gate.
    gbk = parse_workbook("g.csv", "姓名\n张三\n".encode("gbk")).sheets[0]
    assert gbk.columns == ["姓名"]
    assert gbk.rows == [["张三"]]


def test_csv_tab_and_semicolon_delimiters():
    tab = parse_workbook("t.csv", b"a\tb\tc\n1\t2\t3\n").sheets[0]
    assert tab.columns == ["a", "b", "c"]
    assert tab.rows == [["1", "2", "3"]]

    semi = parse_workbook("s.csv", b"a;b;c\n1;2;3\n").sheets[0]
    assert semi.columns == ["a", "b", "c"]


def test_csv_header_offset_and_blank_lines():
    content = b"Title Line\n\nname,city\nAnn,Beijing\n\nBob,Shanghai\n"
    sheet = parse_workbook("offset.csv", content).sheets[0]
    assert sheet.header_row_number == 3
    assert sheet.rows == [["Ann", "Beijing"], ["Bob", "Shanghai"]]


def test_too_many_rows_rejected():
    header = "a,b\n"
    content = header.encode() + b"1,2\n" * (MAX_DATA_ROWS + 1)
    with pytest.raises(ParseError, match="data rows"):
        parse_workbook("big.csv", content)


# ---------- file type / rejection ----------


def test_zip_magic_beats_extension_and_garbage_is_rejected(tmp_path):
    real_xlsx = make_xlsx_bytes({"s": [["a"], ["b"]]})
    renamed = parse_workbook("actually.xlsx.csv", real_xlsx)
    assert renamed.file_type == "xlsx"

    with pytest.raises(ParseError, match="empty"):
        parse_workbook("empty.csv", b"")

    with pytest.raises(ParseError, match=r"\.xls"):
        parse_workbook("legacy.xls", b"anything")

    with pytest.raises(ParseError, match="not a valid XLSX"):
        parse_workbook("broken.xlsx", b"PK\x03\x04garbage-not-really-zip")

    with pytest.raises(ParseError, match="not an XLSX container"):
        parse_workbook("lying.xlsx", b"plain,text\n1,2\n")


# ---------- API ----------


def test_parse_endpoint_accepts_xlsx_and_csv():
    xlsx = make_xlsx_bytes({"客户": [["名"], ["张三"]]})
    xlsx_media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("sales.xlsx", xlsx, xlsx_media)},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["file_type"] == "xlsx"
    assert body["sheets"][0]["columns"] == ["名"]

    csv_response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("c.csv", "名\n张三\n".encode("gbk"), "text/csv")},
    )
    assert csv_response.status_code == 200
    assert csv_response.json()["sheets"][0]["rows"] == [["张三"]]


def test_parse_endpoint_rejects_bad_file_with_422():
    response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("legacy.xls", b"nope", "application/octet-stream")},
    )
    assert response.status_code == 422
    assert ".xls" in response.json()["detail"]


def test_parse_endpoint_rejects_oversize_with_413(monkeypatch):
    monkeypatch.setattr(ingestion_route, "MAX_UPLOAD_BYTES", 10)
    response = client.post(
        "/api/v1/ingestion/parse",
        files={"file": ("big.csv", b"a" * 100, "text/csv")},
    )
    assert response.status_code == 413
